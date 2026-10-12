#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate three editable Word templates for the drug TK/TD individual assignment."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Calibri'
INK = RGBColor(0x1E, 0x20, 0x19)
GREEN = RGBColor(0x2A, 0x3F, 0x2D)
AMBER = RGBColor(0x8F, 0x5D, 0x14)
SOFT = RGBColor(0x5A, 0x5D, 0x51)


def set_base_font(doc):
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(11)
    style.font.color.rgb = INK
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT)


def set_margins(doc, cm=2.0):
    for section in doc.sections:
        section.top_margin = Cm(cm)
        section.bottom_margin = Cm(cm)
        section.left_margin = Cm(cm + 0.2)
        section.right_margin = Cm(cm + 0.2)


def eyebrow(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = SOFT
    p.paragraph_format.space_after = Pt(2)
    return p


def title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = GREEN
    p.paragraph_format.space_after = Pt(14)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = GREEN
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = AMBER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), 'B9791E')
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def body(doc, text, italic=False, size=11, color=None, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(8)
    return p


def fill_line(doc, label, width=45):
    p = doc.add_paragraph()
    r1 = p.add_run(label + ': ')
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run('_' * width)
    r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(10)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    return p


def checkbox(doc, text):
    p = doc.add_paragraph()
    r1 = p.add_run('☐  ')
    r1.font.size = Pt(13)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(8)
    return p


def write_lines(doc, n, note=None):
    if note:
        p = doc.add_paragraph()
        r = p.add_run(note)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = SOFT
        p.paragraph_format.space_after = Pt(4)
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'D6D6C6')
        pbdr.append(bottom)
        pPr.append(pbdr)


def spacer(doc, pt=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pt)


def section_gap(doc, blank_pt=180):
    """Blank space for typing an answer, closed off with a single thin rule —
    replaces the old repeated handwriting lines (doc is meant to be typed)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(blank_pt)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'D6D6C6')
    pbdr.append(bottom)
    pPr.append(pbdr)


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


# ============================================================
# DOKUMENT 1 — Individuální zadání
# ============================================================
doc1 = Document()
set_base_font(doc1)
set_margins(doc1)

eyebrow(doc1, 'Úvod do toxikologie 2026 · Individuální/skupinový úkol')
title(doc1, 'Toxikokinetika a toxikodynamika aktuálně užívané drogy (užívaných drog)')
fill_line(doc1, 'Číslo studujícího', 15)
fill_line(doc1, 'Přidělená látka', 45)
spacer(doc1, 4)
body(doc1, 'Zpracujte samostatně/ve skupině níže uvedené okruhy pro svou přidělenou látku/vaše přidělené látky. Využijte znalosti a koncepty probrané v přednáškách 1–4 a odkazujte se na ně přímo (např. „viz Henderson–Hasselbalchova rovnice, přednáška 3“). Doporučený rozsah: 2–3 normostrany textu (bez obrázků a seznamu zdrojů) na jednu látku.', italic=True, size=10, color=SOFT)

h1(doc1, '1. Identifikace látky')
bullet(doc1, 'Název (čj i aj), běžné slangové názvy')
bullet(doc1, 'Chemická třída a strukturní vzorec, CAS číslo')
bullet(doc1, 'Legální status v ČR (návyková látka dle zákona, léčivo, NPS…)')
section_gap(doc1)

h1(doc1, '2. Toxikokinetika')
bullet(doc1, 'Cesta vstupu typická pro rekreační/zneužívající užití (p.o., inhalace/kouření, intranazálně, injekčně…) a co to znamená pro rychlost nástupu účinku')
bullet(doc1, 'Ionizace a pKa — jak ovlivňuje vstřebávání')
bullet(doc1, 'Distribuce — lipofilita, vazba na plazmatické proteiny, průchod hematoencefalickou bariérou (klíčové u látek působících na CNS), případný depot')
bullet(doc1, 'Biotransformace — které enzymy (CYP, jiné), vzniká aktivní/toxický metabolit? Jde o detoxikaci, nebo bioaktivaci?')
bullet(doc1, 'Eliminace — kinetika (řád), přibližný biologický poločas, význam pro dobu detekovatelnosti')
section_gap(doc1)

doc1.add_page_break()

h1(doc1, '3. Toxikodynamika')
bullet(doc1, 'Molekulární cíl — receptor, iontový kanál, enzym, transportér')
bullet(doc1, 'Typ interakce — agonista, antagonista, inhibitor zpětného vychytávání…')
bullet(doc1, 'Mechanismus akutní toxicity a předávkování — co se děje na buněčné/orgánové úrovni')
bullet(doc1, 'Tolerance a závislost — mechanismus (např. down-regulace receptorů), abstinenční příznaky')
section_gap(doc1)

h1(doc1, '4. Reálný rozměr')
bullet(doc1, 'Konkrétní kazuistika nebo epidemiologická data (reálný případ otravy/úmrtí, pokud dohledatelný)')
bullet(doc1, 'Významné lékové/drogové interakce (často přes sdílený CYP enzym)')
bullet(doc1, 'Antidotum nebo léčba, pokud existuje')
section_gap(doc1)

h1(doc1, '5. Použité zdroje')
body(doc1, 'Uveďte minimálně 4 odborné zdroje (ne Wikipedii). Postupujte podle doporučeného schématu práce s literaturou na stránce úkolu — chemická data (PubChem/DrugBank) → toxikologický kontext (TIS/PubMed) → epidemiologický rozměr (EMCDDA/NMS ČR) → případně s kritickým odstupem (Erowid).', italic=True, size=10, color=SOFT)
section_gap(doc1)

doc1.save('Sablona_individualni_zadani.docx')
print('written Sablona_individualni_zadani.docx')


# ============================================================
# DOKUMENT 2 — Pracovní list, krok B
# ============================================================
doc2 = Document()
set_base_font(doc2)
set_margins(doc2)

eyebrow(doc2, 'Úvod do toxikologie 2026 · Seminář 8. prosince')
title(doc2, 'Pracovní list — krok B: vzájemná výuka ve skupině')
fill_line(doc2, 'Skupina č.', 8)
fill_line(doc2, 'Moje číslo a látka', 40)
spacer(doc2, 4)
body(doc2, 'Postupně vysvětlete ostatním členům skupiny svou látku (bez internetu, jen z poznámek). Zatímco posloucháte spolužáky, vyplňte tabulku — pro každého jednoho z vašich čtyř spolužáků jeden řádek. Nezapomeňte položit alespoň jednu doplňující otázku ke každé látce.', italic=True, size=10, color=SOFT)
body(doc2, 'Tento list po semináři odevzdejte (foto/sken) do své složky v Teams — hodnotí se individuálně.', bold=True, size=10)
spacer(doc2, 6)

table = doc2.add_table(rows=5, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True
headers = ['Látka spolužáka', 'Hlavní mechanismus (2–3 věty vlastními slovy)', 'Nejzajímavější riziko / fakt', 'Otázka, kterou jsem položil/a']
widths = [Cm(3.0), Cm(4.6), Cm(4.6), Cm(3.6)]
for i, htext in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    shade_cell(cell, 'EAF1EA')
    p = cell.paragraphs[0]
    r = p.add_run(htext)
    r.bold = True
    r.font.size = Pt(10)
for row_idx in range(1, 5):
    for col_idx in range(4):
        cell = table.rows[row_idx].cells[col_idx]
        cell.width = widths[col_idx]
        p = cell.paragraphs[0]
        p.add_run(' ')
        for _ in range(2):
            cell.add_paragraph(' ')

doc2.save('Pracovni_list_krok_B.docx')
print('written Pracovni_list_krok_B.docx')


# ============================================================
# DOKUMENT 3 — Syntézní karta, krok C
# ============================================================
doc3 = Document()
set_base_font(doc3)
set_margins(doc3)

eyebrow(doc3, 'Úvod do toxikologie 2026 · Seminář 8. prosince')
title(doc3, 'Skupinová syntéza — krok C')
fill_line(doc3, 'Skupina č.', 8)
fill_line(doc3, 'Členové (čísla)', 40)
body(doc3, 'Tuto kartu po semináři odevzdejte do Teams (jedna za celou skupinu) — hodnotí se sdíleně pro všechny členy skupiny.', bold=True, size=10)
spacer(doc3, 4)

h2(doc3, 'Zvolte úhel syntézy (zaškrtněte jednu možnost)')
checkbox(doc3, '1. Časová osa účinku — porovnání nástupu, vrcholu a odeznění účinku všech 5 látek')
checkbox(doc3, '2. Mapa cílů v těle — kde látky zasahují (receptory/enzymy/kanály), co mají společné a odlišné')
checkbox(doc3, '3. Žebříček rizika — seřazení podle letality/bezpečnostního okna s odůvodněním')
checkbox(doc3, '4. Detoxikace vs. bioaktivace — jak si tělo s látkami poradí, kde vzniká nebezpečný metabolit')
checkbox(doc3, '5. „Co se stane, když…“ — fiktivní kazuistika kombinace dvou a více látek ze skupiny')
checkbox(doc3, '6. Vlastní nápad (musí použít min. 2 koncepty z přednášek a zahrnout všech 5 látek)')

h2(doc3, 'Odůvodnění volby')
section_gap(doc3, 120)

h2(doc3, 'Hlavní poznatek pro plénum')
body(doc3, 'Co chcete sdělit třídě ve 4minutové prezentaci — max. 4–5 bodů.', italic=True, size=10, color=SOFT)
section_gap(doc3, 180)

h2(doc3, 'Podklady použité ve skupině')
section_gap(doc3, 120)

doc3.save('Synteticka_karta_krok_C.docx')
print('written Synteticka_karta_krok_C.docx')
